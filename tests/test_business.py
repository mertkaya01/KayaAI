"""Adım 10: iş modülleri — okuyucular, analiz (sayıları kod hesaplar), PDF özeti, fatura, rapor (md/xlsx/docx/pdf)."""

from __future__ import annotations

import io
import json
from datetime import date
from pathlib import Path

import pytest

from kayaai.business import ReadError, read_pdf, read_table, render, run_mode, save
from kayaai.business.analysis import AnalysisPlan, execute, unverified_numbers, validate_plan
from kayaai.business.documents import Invoice, InvoiceLine, check_invoice
from kayaai.business.numbers import fmt, parse_number
from kayaai.business.report import Report, ReportTable, from_markdown, from_result
from kayaai.config import load_settings
from kayaai.runtime import build_runtime

from .conftest import good_answer

ROWS = [
    ("Ocak", "2026-01-15", "Ankara", "Kalem", 120.0),
    ("Ocak", "2026-01-20", "İzmir", "Defter", 80.0),
    ("Şubat", "2026-02-03", "Ankara", "Defter", 200.0),
    ("Şubat", "2026-02-10", "İstanbul", "Kalem", 50.5),
    ("Mart", "2026-03-01", "İzmir", "Kalem", 99.5),
]


def make_xlsx(rows=ROWS, text_numbers: bool = False) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Satışlar"
    ws.append(["Ay", "Tarih", "Bölge", "Ürün", "Tutar"])
    for ay, d, bolge, urun, tutar in rows:
        ws.append([ay, date.fromisoformat(d), bolge, urun, fmt(tutar) + " TL" if text_numbers else tutar])
    wb.create_sheet("Boş")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_pdf(pages: list[str], encrypt: bool = False, text: bool = True) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    from kayaai.business.report import _register_font

    # Sistem fontu (Windows: Segoe UI/Arial, Linux: DejaVu) — v0.10.1: önceden Linux yolu sabitti, Windows'ta düşüyordu
    font, _ = _register_font()
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for page in pages:
        if text:
            y = 800
            for line in page.split("\n"):
                c.setFont(font, 11)
                c.drawString(50, y, line)
                y -= 16
        else:
            c.rect(50, 50, 400, 600, fill=1)
        c.showPage()
    c.save()
    data = buf.getvalue()
    if encrypt:
        from pypdf import PdfReader, PdfWriter

        w = PdfWriter()
        for p in PdfReader(io.BytesIO(data)).pages:
            w.add_page(p)
        w.encrypt("gizli123")
        out = io.BytesIO()
        w.write(out)
        data = out.getvalue()
    return data


INVOICE_TEXT = ("FATURA\nSatıcı: Kaya Kırtasiye Ltd. Şti.\nAlıcı: Atlas Yazılım A.Ş.\nFatura No: KYA2026000123\n"
                "Fatura Tarihi: 15.01.2026\nKalem: Defter 10 adet x 25,00 = 250,00 TL\nKalem: Kalem 50 adet x 5,00 = 250,00 TL\n"
                "Mal Hizmet Toplam Tutarı: 500,00 TL\nHesaplanan KDV (%20): 100,00 TL\nÖdenecek Tutar: 600,00 TL")


@pytest.fixture
def env(fake_ollama, tmp_path: Path, monkeypatch):
    monkeypatch.setenv("KAYAAI_OLLAMA_URL", fake_ollama.url)
    monkeypatch.setenv("KAYAAI_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("KAYAAI_FILES_DIR", str(tmp_path / "files"))
    s = load_settings()
    s.ollama.max_retries = 0
    return s


@pytest.fixture
def rt(env):
    r = build_runtime(env)
    yield r
    r.close()


def by_title(**replies):
    """Şema başlığına göre sırayla cevap veren sahte model."""
    queues = {k: list(v) for k, v in replies.items()}
    seen: dict[str, list[str]] = {}

    def reply(body):
        title = (body.get("format") or {}).get("title", "?")
        seen.setdefault(title, []).append(body["messages"][-1]["content"])
        q = queues.get(title, [])
        return q.pop(0) if q else "{}"
    reply.seen = seen  # type: ignore[attr-defined]
    return reply


# ---------------------------------------------------------------- sayılar ve okuyucular

@pytest.mark.parametrize("raw,expected", [
    ("1.234,56 ₺", 1234.56), ("1,234.56", 1234.56), ("12,5", 12.5), ("1.234", 1234.0), ("%18", 18.0),
    ("(250)", -250.0), ("TL 2.500", 2500.0), ("abc", None), ("", None), (42, 42.0), ("1.234.567", 1234567.0),
])
def test_parse_number(raw, expected):
    assert parse_number(raw) == expected


def test_fmt_turkish():
    assert fmt(1234.5) == "1.234,50" and fmt(1200.0) == "1.200" and fmt(None) == ""


def test_read_xlsx_types_and_turkish_numbers():
    [t] = read_table(make_xlsx(text_numbers=True), "satis.xlsx")  # boş sayfa atlanır
    assert t.name == "Satışlar" and t.columns == ["Ay", "Tarih", "Bölge", "Ürün", "Tutar"]
    assert t.types == ["text", "date", "text", "text", "number"]
    assert t.rows[0][1] == "2026-01-15" and t.rows[3][4] == 50.5


def test_read_csv_semicolon_cp1254():
    data = "Bölge;Tutar\nİzmir;1.250,50\nAnkara;300\n".encode("cp1254")
    [t] = read_table(data, "rapor.csv")
    assert t.columns == ["Bölge", "Tutar"] and t.types == ["text", "number"] and t.rows[0] == ["İzmir", 1250.5]


def test_read_errors():
    with pytest.raises(ReadError, match="xlsx"):
        read_table(b"x", "eski.xls")
    with pytest.raises(ReadError, match="okunamadı"):
        read_table(b"bozuk", "a.xlsx")
    with pytest.raises(ReadError, match="metin içermiyor"):
        read_pdf(make_pdf(["x"], text=False), "tarama.pdf")
    with pytest.raises(ReadError, match="parola"):
        read_pdf(make_pdf(["Gizli sözleşme metni burada yazıyor."], encrypt=True), "gizli.pdf")


def test_read_pdf_turkish():
    doc = read_pdf(make_pdf(["Şubat ayı İzmir satışları çok iyi geçti.", "İkinci sayfa: öneriler ve ölçümler."]), "r.pdf")
    assert len(doc.pages) == 2 and "Şubat" in doc.pages[0] and "[s.2]" in doc.text


# ---------------------------------------------------------------- analiz

def plan(**op) -> AnalysisPlan:
    return AnalysisPlan.model_validate({"operations": [{"title": "Sonuç", **op}]})


def test_validate_plan_catches_bad_columns_and_types():
    tables = read_table(make_xlsx(), "s.xlsx")
    errs = validate_plan(plan(metrics=[{"agg": "sum", "column": "Fiyat"}]), tables)
    assert any("bilinmeyen sütun" in e for e in errs)
    errs = validate_plan(plan(metrics=[{"agg": "sum", "column": "Bölge"}]), tables)
    assert any("sayısal değil" in e for e in errs)
    errs = validate_plan(plan(group_by=["Bölge"], period="month", metrics=[{"agg": "count"}]), tables)
    assert any("period" in e for e in errs)
    assert validate_plan(plan(group_by=["bölge"], metrics=[{"agg": "sum", "column": "tutar"}]), tables) == []


def test_execute_group_share_period_filter():
    tables = read_table(make_xlsx(), "s.xlsx")
    [r] = execute(plan(group_by=["Bölge"], metrics=[{"agg": "sum", "column": "Tutar"}]), tables)
    assert r.columns == ["Bölge", "Toplam (Tutar)", "Pay (%)"]
    assert r.rows[0] == ["Ankara", 320.0, 58.18] and sum(x[1] for x in r.rows) == 550.0
    [m] = execute(plan(group_by=["Tarih"], period="month", sort="group_asc",
                       metrics=[{"agg": "sum", "column": "Tutar"}, {"agg": "count"}]), tables)
    assert [row[0] for row in m.rows] == ["2026-01", "2026-02", "2026-03"] and m.rows[1][1:3] == [250.5, 2.0]
    [f] = execute(plan(filters=[{"column": "Ürün", "op": "=", "value": "kalem"}, {"column": "Tutar", "op": ">", "value": "60"}],
                       metrics=[{"agg": "sum", "column": "Tutar"}, {"agg": "mean", "column": "Tutar"}]), tables)
    assert f.rows == [[219.5, 109.75]]
    [top] = execute(plan(group_by=["Ürün"], metrics=[{"agg": "max", "column": "Tutar"}], limit=1), tables)
    assert top.rows == [["Defter", 200.0]] and top.total_rows == 2


def test_unverified_numbers():
    allowed = [550.0, 320.0, 58.13, 110.0]
    assert unverified_numbers("Toplam 550 TL; Ankara 320 (%58,13). 3 bölge var.", allowed) == []
    assert unverified_numbers("Ankara payı yaklaşık %58 ve toplam 550,00.", allowed) == []
    assert unverified_numbers("Toplam 777 TL oldu, 2026 yılında.", allowed) == [777.0]


def test_analyze_end_to_end_rejects_invented_numbers(rt, fake_ollama):
    p = json.dumps({"sheet": None, "operations": [{"title": "Bölgelere göre satış", "group_by": ["Bölge"],
                                                   "metrics": [{"agg": "sum", "column": "Tutar"}]}]})
    bad_plan = json.dumps({"operations": [{"title": "x", "metrics": [{"agg": "sum", "column": "Fiyat"}]}]})
    fake_ollama.chat_replies = [by_title(
        AnalysisPlan=[bad_plan, p],
        GeneralAnswer=[good_answer("En çok satış Ankara'da: 999 TL."), good_answer("En çok satış Ankara'da: 320 TL (%58,18).")])] * 10
    tables = read_table(make_xlsx(), "s.xlsx")
    from kayaai.business import analyze
    res = analyze(rt, tables, "Hangi bölge en çok sattı?", "s.xlsx")
    assert res.results[0].rows[0][:2] == ["Ankara", 320.0]
    assert res.narrative == "En çok satış Ankara'da: 320 TL (%58,18)." and res.unverified_numbers == []
    md = res.markdown()
    assert "| Ankara | 320 | 58,18 |" in md and "programla hesaplandı" in md
    reply = fake_ollama.chat_replies[0]
    assert "Fiyat" in reply.seen["AnalysisPlan"][1]  # geçersiz plan geri bildirimle düzeltildi
    assert "999" in reply.seen["GeneralAnswer"][1]   # uydurma sayı geri bildirimi
    assert "Ocak" not in reply.seen["GeneralAnswer"][0] or "result_tables" in reply.seen["GeneralAnswer"][0]


def test_analyze_keeps_warning_when_model_keeps_inventing(rt, fake_ollama):
    p = json.dumps({"operations": [{"title": "Toplam", "metrics": [{"agg": "sum", "column": "Tutar"}]}]})
    fake_ollama.chat_replies = [by_title(AnalysisPlan=[p], GeneralAnswer=[good_answer("Toplam 12.345 TL.")] * 3)] * 10
    res = run_mode(rt, "analyze", make_xlsx(), "s.xlsx", "toplam ne?")
    assert res.unverified_numbers == [12345.0] and any("olmayan sayı" in w for w in res.warnings)
    assert res.results[0].rows == [[550.0]]


def test_analyze_plan_failure_falls_back_to_profile(rt, fake_ollama):
    fake_ollama.chat_replies = [by_title()] * 10
    res = run_mode(rt, "analyze", make_xlsx(), "s.xlsx", "?")
    assert res.results == [] and "Tablo profili" in res.markdown() and res.warnings


# ---------------------------------------------------------------- PDF özeti ve fatura

def test_summarize_short_and_long(rt, fake_ollama):
    fake_ollama.chat_replies = [by_title(GeneralAnswer=[good_answer("- Satışlar arttı [s.1]")])] * 5
    short = run_mode(rt, "summarize", make_pdf(["Satışlar bu ay yüzde on arttı ve hedef aşıldı."]), "k.pdf")
    assert short.summary == "- Satışlar arttı [s.1]" and short.pages == 1
    long_pages = [("Bölüm metni uzun bir cümledir. " * 9).strip()[:240] for _ in range(60)]
    def reply(body):
        q = body["messages"][-1]["content"]
        return good_answer("Genel özet ve maddeler [s.1]" if "bölüm özetleri var" in q else "- Bölüm özeti [s.1]")
    fake_ollama.chat_replies = [reply] * 20
    res = run_mode(rt, "summarize", make_pdf(long_pages), "uzun.pdf", "Ana noktalar neler?")
    assert len(res.sections) >= 2 and res.summary == "Genel özet ve maddeler [s.1]"


def test_check_invoice():
    good = Invoice(seller="Kaya", total=600.0, subtotal=500.0, vat=100.0, vat_rate=20,
                   lines=[InvoiceLine(description="Defter", quantity=10, unit_price=25, amount=250),
                          InvoiceLine(description="Kalem", quantity=50, unit_price=5, amount=250)])
    assert check_invoice(good, INVOICE_TEXT) == []
    bad = good.model_copy(update={"total": 650.0, "lines": good.lines[:1]})
    w = check_invoice(bad, INVOICE_TEXT)
    assert any("Satır toplamı" in x for x in w) and any("uyuşmuyor" in x for x in w) and any("bulunamadı" in x for x in w)


def test_extract_invoice(rt, fake_ollama):
    inv = {"seller": "Kaya Kırtasiye Ltd. Şti.", "buyer": "Atlas Yazılım A.Ş.", "invoice_no": "KYA2026000123",
           "date": "2026-01-15", "due_date": None, "currency": "TRY", "subtotal": 500, "vat_rate": 20, "vat": 100,
           "total": 600, "lines": [{"description": "Defter", "quantity": 10, "unit_price": 25, "amount": 250},
                                   {"description": "Kalem", "quantity": 50, "unit_price": 5, "amount": 250}]}
    fake_ollama.chat_replies = [by_title(Invoice=[json.dumps({**inv, "date": "15.01.2026"}), json.dumps(inv)])] * 5
    res = run_mode(rt, "invoice", make_pdf(INVOICE_TEXT.split("\n")[:1] and [INVOICE_TEXT]), "fatura.pdf")
    assert res.invoice.total == 600 and res.warnings == []
    assert "| Genel toplam | 600 |" in res.markdown() and "tutarlı" in res.markdown()


# ---------------------------------------------------------------- raporlar

def sample_report() -> Report:
    return Report(title="Şubat Satış Raporu", summary="- İzmir öne çıktı\n- Ölçümler **arttı**",
                  key_figures=[("Toplam", "550")],
                  tables=[ReportTable(title="Bölgeler", columns=["Bölge", "Tutar"], rows=[["İzmir", 179.5], ["Ankara", 320.0]])],
                  sources=["satis.xlsx"], notes=["Sayılar programla hesaplandı."])


def test_font_search_prefers_windows_fonts():
    from kayaai.business.report import FONT_CANDIDATES
    assert FONT_CANDIDATES[0][0].startswith("C:/Windows/Fonts/") and FONT_CANDIDATES[1][0].startswith("C:/Windows/Fonts/")


def test_four_formats_roundtrip():
    r = sample_report()
    md = render(r, "md").decode("utf-8")
    assert md.startswith("# Şubat Satış Raporu") and "| İzmir | 179,50 |" in md
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(render(r, "xlsx")))
    assert wb.sheetnames == ["Özet", "Bölgeler"] and wb["Bölgeler"]["B3"].value == 320.0
    assert wb["Özet"]["A1"].value == "Şubat Satış Raporu"
    from docx import Document
    doc = Document(io.BytesIO(render(r, "docx")))
    texts = [p.text for p in doc.paragraphs]
    assert "Şubat Satış Raporu" in texts and "İzmir öne çıktı" in texts and doc.tables[1].cell(1, 0).text == "İzmir"
    from pypdf import PdfReader
    pdf_text = PdfReader(io.BytesIO(render(r, "pdf"))).pages[0].extract_text()
    assert "Şubat Satış Raporu" in pdf_text and "İzmir" in pdf_text and "Ölçümler" in pdf_text


def test_from_result_and_markdown():
    rep = from_result({"kind": "analysis", "source": "s.xlsx", "sheet": "Satışlar", "narrative": "Ankara önde.",
                       "profile": {"rows": 5, "columns": [{"name": "Tutar", "type": "number", "sum": 550.0}]},
                       "results": [{"title": "Bölgeler", "columns": ["Bölge", "Toplam"], "rows": [["Ankara", 320.0]],
                                    "total_rows": 1}], "warnings": []})
    assert rep.tables[0].rows == [["Ankara", 320.0]] and ("Tutar toplamı", "550") in rep.key_figures
    m = from_markdown("Giriş paragrafı.\n\n**Bölgeler**\n\n| Bölge | Tutar |\n|---|---|\n| İzmir | 10 |\n\nSon söz.", "Başlık")
    assert m.tables[0].title == "Bölgeler" and m.tables[0].rows == [["İzmir", "10"]] and "Son söz." in m.summary


def test_save_new_is_r1_overwrite_needs_approval(rt, env):
    r = sample_report()
    first = save(rt, r, "xlsx", filename="subat")
    path = env.data_dir.parent / "files" / "raporlar" / "subat.xlsx"
    assert first.status == "done" and first.risk.value == "R1" and path.is_file()
    second = save(rt, r, "pdf", filename="subat")
    assert second.status == "done"  # farklı uzantı = yeni dosya
    third = save(rt, r, "xlsx", filename="subat")
    assert third.status == "awaiting_approval" and third.risk.value == "R2"
    a = rt.permissions.approve(third.approval_id, decided_by="test")
    assert a.status == "approved" and a.result["overwritten"] and Path(a.result["backup"]).is_file()


# ---------------------------------------------------------------- sohbet / API / CLI / MCP / belge araması

PLAN_BY_REGION = json.dumps({"operations": [{"title": "Bölgelere göre satış", "group_by": ["Bölge"],
                                             "metrics": [{"agg": "sum", "column": "Tutar"}]}]})


def test_chat_table_attachment_is_analyzed_and_reported(rt, fake_ollama, env):
    import base64

    from kayaai.assistant import ChatBook
    book = ChatBook(rt)
    conv = book.create()
    att = book.upload_binary("satis.xlsx", base64.b64encode(make_xlsx()).decode())
    assert att.kind == "table" and "Ankara" in book.attachment(att.id)[1]
    fake_ollama.chat_replies = [by_title(AnalysisPlan=[PLAN_BY_REGION], GeneralAnswer=[good_answer("Ankara 320 ile önde.")])] * 10
    user, reply, run = book.send(conv.id, "Hangi bölge en çok sattı?", attachment_ids=[att.id])
    run()
    msg = book.view(conv.id).messages[-1]
    assert msg.status == "done" and msg.agent == "data" and "| Ankara | 320 |" in msg.content
    assert msg.meta["business"]["kind"] == "analysis"
    for fmt_ in ("xlsx", "pdf", "docx", "md"):
        res = book.report(conv.id, msg.id, fmt_)
        assert res.status == "done" and Path(res.result["path"]).suffix == f".{fmt_}"
    files = sorted(p.suffix for p in (env.data_dir.parent / "files" / "raporlar").iterdir())
    assert files == [".docx", ".md", ".pdf", ".xlsx"]
    from kayaai.assistant.chat import ChatError
    with pytest.raises(ChatError):
        book.report(conv.id, user.id, "pdf")  # kullanıcı mesajından rapor olmaz


def test_chat_pdf_modes_and_free_text_report(rt, fake_ollama):
    import base64

    from kayaai.assistant import ChatBook
    book = ChatBook(rt)
    conv = book.create()
    att = book.upload_binary("fatura.pdf", base64.b64encode(make_pdf([INVOICE_TEXT])).decode())
    assert att.kind == "pdf"
    inv = {"seller": "Kaya Kırtasiye", "invoice_no": "KYA2026000123", "date": "2026-01-15", "subtotal": 500,
           "vat": 100, "vat_rate": 20, "total": 600, "lines": []}
    fake_ollama.chat_replies = [by_title(Invoice=[json.dumps(inv)])] * 5
    _, _, run = book.send(conv.id, "Faturayı oku", attachment_ids=[att.id], mode="invoice")
    run()
    msg = book.view(conv.id).messages[-1]
    assert msg.status == "done" and "| Genel toplam | 600 |" in msg.content and msg.warnings == []
    from kayaai.assistant.chat import ChatError
    with pytest.raises(ChatError, match="Excel/CSV"):
        book.send(conv.id, "analiz et", attachment_ids=[att.id], mode="analyze")
    # PDF eki + normal soru → metin bağlamı (sayfa işaretli)
    seen = []
    fake_ollama.chat_replies = [lambda b: (seen.append(b["messages"][-1]["content"]),
                                           json.dumps({"mode": "direct", "agent": "general", "steps": [], "needs_claude": False,
                                                       "claude_reason": "", "key_topics": []})
                                           if (b.get("format") or {}).get("title") == "RoutePlan"
                                           else good_answer("Fatura toplamı 600 TL.\n\n| Kalem | Tutar |\n|---|---|\n| Defter | 250 |"))[1]] * 10
    _, _, run2 = book.send(conv.id, "Bu faturanın toplamı ne?", attachment_ids=[att.id])
    run2()
    free = book.view(conv.id).messages[-1]
    assert any("[s.1]" in s and "Ödenecek Tutar" in s for s in seen)
    res = book.report(conv.id, free.id, "xlsx")
    assert res.status == "done"


def test_binary_upload_rejections(rt):
    import base64

    from kayaai.assistant import ChatBook
    from kayaai.assistant.chat import ChatError
    book = ChatBook(rt)
    with pytest.raises(ChatError, match="türü"):
        book.upload_binary("virus.exe", base64.b64encode(b"MZ").decode())
    with pytest.raises(ChatError, match="bozuk"):
        book.upload_binary("a.pdf", "%%%")
    with pytest.raises(ChatError, match="metin içermiyor"):
        book.upload_binary("tarama.pdf", base64.b64encode(make_pdf(["x"], text=False)).decode())


def test_document_index_reads_pdf_and_xlsx(rt, env):
    files = env.data_dir.parent / "files"
    files.mkdir(parents=True, exist_ok=True)
    (files / "rapor.pdf").write_bytes(make_pdf(["Giriş sayfası.", "Atlas projesinin bütçesi 120.000 TL olarak onaylandı."]))
    (files / "satis.xlsx").write_bytes(make_xlsx())
    (files / "tarama.pdf").write_bytes(make_pdf(["x"], text=False))
    rt.settings.memory.min_score_docs = 0.1
    rep = rt.memory.docs.index()
    names = {Path(d.path).name for d in rt.memory.docs.documents()}
    assert names == {"rapor.pdf", "satis.xlsx"} and rep.skipped.get("okunamadı (PDF/Excel)") == 1
    hit = rt.memory.recall("Atlas bütçesi ne kadar onaylandı?").documents[0]
    assert hit["name"] == "rapor.pdf" and hit["loc"] in ("s.1", "s.2")


def test_api_business_and_report(env, fake_ollama):
    import base64

    from fastapi.testclient import TestClient

    from kayaai.api import create_app
    rt_ = build_runtime(env)
    app = create_app(env, runtime=rt_)
    files = env.data_dir.parent / "files"
    files.mkdir(parents=True, exist_ok=True)
    (files / "s.xlsx").write_bytes(make_xlsx())
    fake_ollama.chat_replies = [by_title(AnalysisPlan=[PLAN_BY_REGION] * 3, GeneralAnswer=[good_answer("Ankara önde: 320.")] * 3)] * 20
    with TestClient(app) as c:
        c.headers["Authorization"] = f"Bearer {app.state.token}"
        r = c.post("/api/business/run", json={"path": "s.xlsx", "mode": "analyze", "question": "bölgeler", "report": "pdf"})
        assert r.status_code == 200 and r.json()["result"]["results"][0]["rows"][0][0] == "Ankara"
        assert r.json()["report"]["status"] == "done" and r.json()["report"]["result"]["path"].endswith(".pdf")
        assert c.post("/api/business/run", json={"path": "/etc/passwd", "mode": "analyze"}).status_code == 403
        assert c.post("/api/business/run", json={"path": "yok.xlsx", "mode": "analyze"}).status_code == 404
        att = c.post("/api/attachments", json={"name": "s.xlsx", "data_b64": base64.b64encode(make_xlsx()).decode()}).json()
        assert att["kind"] == "table"
        conv = c.post("/api/chat", json={}).json()
        c.post(f"/api/chat/{conv['id']}/messages", json={"text": "analiz et", "attachments": [att["id"]]})
        import time
        for _ in range(200):
            msgs = c.get(f"/api/chat/{conv['id']}").json()["messages"]
            if msgs[-1]["status"] in ("done", "failed"):
                break
            time.sleep(0.05)
        assert msgs[-1]["status"] == "done"
        rep = c.post(f"/api/chat/{conv['id']}/messages/{msgs[-1]['id']}/report", json={"format": "docx"})
        assert rep.status_code == 200 and rep.json()["status"] == "done"
        assert c.post(f"/api/chat/{conv['id']}/messages/{msgs[-1]['id']}/report", json={"format": "exe"}).status_code == 422
    rt_.close()


def test_cli_analyze_with_report(env, fake_ollama, tmp_path, capsys):
    from kayaai import cli
    f = tmp_path / "masaustu.xlsx"
    f.write_bytes(make_xlsx())
    fake_ollama.chat_replies = [by_title(AnalysisPlan=[PLAN_BY_REGION], GeneralAnswer=[good_answer("Ankara önde: 320.")])] * 10
    assert cli.main(["analyze", str(f), "hangi", "bölge", "--report", "xlsx"]) == 0
    out = capsys.readouterr().out
    assert "| Ankara | 320 |" in out and "[RAPOR]" in out and ".xlsx" in out
    assert cli.main(["summarize", str(tmp_path / "yok.pdf")]) == 1


def test_mcp_business_tools(env, fake_ollama):
    import asyncio

    from mcp import Client

    from kayaai.mcp_server import create_server
    files = env.data_dir.parent / "files"
    files.mkdir(parents=True, exist_ok=True)
    (files / "s.xlsx").write_bytes(make_xlsx())
    fake_ollama.chat_replies = [by_title(AnalysisPlan=[PLAN_BY_REGION], GeneralAnswer=[good_answer("Ankara önde: 320.")])] * 10

    async def go():
        async with Client(create_server(env)) as c:
            ok = await c.call_tool("kayaai_analyze_table", {"path": "s.xlsx", "question": "bölgeler"})
            bad = await c.call_tool("kayaai_read_pdf", {"path": "/etc/passwd"})
            return ok, bad

    ok, bad = asyncio.run(go())
    assert "| Ankara | 320 |" in ok.structured_content["markdown"] and bad.is_error

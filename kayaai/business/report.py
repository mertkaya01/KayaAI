"""İş raporu (Adım 10, kabul #12). Sahip kararı: Excel, Markdown, Word ve PDF.

Tek rapor modeli → dört yazıcı. Kayıt `~/kayaAI-files/raporlar/` altına, mevcut dosya araçları ve izin motoru ile:
yeni dosya R1 (otomatik, kayda geçer), var olanın üzerine yazma R2 (sahip onayı + yedek).
PDF'te Türkçe karakterler için sistem fontu kullanılır (repoya font eklenmez).
"""

from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING, Any, Literal

from pydantic import BaseModel, Field

from .numbers import fmt

if TYPE_CHECKING:
    from ..core.permissions import SubmitResult
    from ..runtime import Runtime

Format = Literal["md", "xlsx", "docx", "pdf"]
FORMATS: tuple[Format, ...] = ("md", "xlsx", "docx", "pdf")
REPORT_DIR = "raporlar"


class ReportError(ValueError):
    pass


class ReportTable(BaseModel):
    title: str
    columns: list[str]
    rows: list[list[Any]]


class Report(BaseModel):
    title: str
    created_at: datetime = Field(default_factory=datetime.now)
    summary: str = ""                       # düz metin / hafif Markdown (madde işaretleri, **kalın**)
    key_figures: list[tuple[str, str]] = Field(default_factory=list)
    tables: list[ReportTable] = Field(default_factory=list)
    sources: list[str] = Field(default_factory=list)
    notes: list[str] = Field(default_factory=list)  # uyarılar / doğrulama notları


# ---------------------------------------------------------------------------- kaynaklardan rapor


def _cell(v: Any) -> str:
    return fmt(v) if isinstance(v, float) else ("" if v is None else str(v))


def from_result(result: dict[str, Any], title: str | None = None) -> Report:
    """Analiz / özet / fatura sonucundan (mesajda saklanan JSON) rapor."""
    kind = result.get("kind")
    if kind == "analysis":
        rep = Report(title=title or f"Veri analizi — {Path(result['source']).stem}",
                     summary=result.get("narrative", ""), sources=[f"{result['source']} (sayfa: {result.get('sheet')})"])
        prof = result.get("profile", {})
        rep.key_figures.append(("Satır sayısı", fmt(float(prof.get("rows", 0)))))
        for c in prof.get("columns", [])[:6]:
            if c.get("type") == "number" and "sum" in c:
                rep.key_figures.append((f"{c['name']} toplamı", fmt(c["sum"])))
        rep.tables = [ReportTable(title=r["title"], columns=r["columns"], rows=r["rows"]) for r in result.get("results", [])]
        rep.notes = list(result.get("warnings", [])) + ["Tüm sayılar kayaAI tarafından dosyadan programla hesaplandı."]
        return rep
    if kind == "summary":
        return Report(title=title or f"Belge özeti — {Path(result['source']).stem}", summary=result.get("summary", ""),
                      sources=[f"{result['source']} ({result.get('pages')} sayfa)"], notes=list(result.get("warnings", [])))
    if kind == "invoice":
        inv = result.get("invoice") or {}
        rep = Report(title=title or f"Fatura — {inv.get('invoice_no') or Path(result['source']).stem}",
                     sources=[result["source"]], notes=list(result.get("warnings", [])) or ["Tutarlar kodla denetlendi: tutarlı."])
        rep.key_figures = [(k, v) for k, v in (
            ("Satıcı", inv.get("seller")), ("Alıcı", inv.get("buyer")), ("Fatura no", inv.get("invoice_no")),
            ("Tarih", inv.get("date")), ("Vade", inv.get("due_date")), ("Ara toplam", fmt(inv.get("subtotal"))),
            ("KDV", fmt(inv.get("vat"))), ("Genel toplam", fmt(inv.get("total"))), ("Para birimi", inv.get("currency")))
            if v]
        if inv.get("lines"):
            rep.tables = [ReportTable(title="Kalemler", columns=["Açıklama", "Miktar", "Birim fiyat", "Tutar"],
                                      rows=[[ln.get("description"), ln.get("quantity"), ln.get("unit_price"), ln.get("amount")]
                                            for ln in inv["lines"]])]
        return rep
    raise ReportError("Bu sonuçtan rapor üretilemiyor")


_MD_TABLE = re.compile(r"^\s*\|.*\|\s*$")


def from_markdown(text: str, title: str, sources: list[str] | None = None) -> Report:
    """Serbest bir sohbet cevabından rapor: Markdown tabloları tabloya, geri kalanı özete."""
    rep = Report(title=title, sources=sources or [])
    lines = text.replace("\r\n", "\n").split("\n")
    summary: list[str] = []
    i = 0
    pending_title = ""
    while i < len(lines):
        if _MD_TABLE.match(lines[i]) and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            cols = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and _MD_TABLE.match(lines[i]):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            rep.tables.append(ReportTable(title=pending_title or f"Tablo {len(rep.tables) + 1}", columns=cols, rows=rows))
            pending_title = ""
            continue
        line = lines[i]
        m = re.match(r"^\s*\*\*(.+)\*\*\s*$", line)
        if m:
            pending_title = m.group(1)
        else:
            summary.append(line)
        i += 1
    rep.summary = "\n".join(summary).strip()
    return rep


# ---------------------------------------------------------------------------- yazıcılar


def _plain(text: str) -> str:
    return re.sub(r"[*_`]+", "", text)


def to_markdown(r: Report) -> bytes:
    out = [f"# {r.title}", f"_{r.created_at:%d.%m.%Y %H:%M} — kayaAI_", ""]
    if r.summary:
        out += ["## Özet", "", r.summary.strip(), ""]
    if r.key_figures:
        out += ["## Anahtar rakamlar", "", "| | |", "|---|---|"] + [f"| {k} | {v} |" for k, v in r.key_figures] + [""]
    for t in r.tables:
        out += [f"## {t.title}", "", "| " + " | ".join(t.columns) + " |", "|" + "---|" * len(t.columns)]
        out += ["| " + " | ".join(_cell(v) for v in row) + " |" for row in t.rows] + [""]
    if r.notes:
        out += ["## Notlar", ""] + [f"- {n}" for n in r.notes] + [""]
    if r.sources:
        out += ["## Kaynaklar", ""] + [f"- {s}" for s in r.sources] + [""]
    return "\n".join(out).encode("utf-8")


def to_xlsx(r: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    bold, head_fill = Font(bold=True), PatternFill("solid", fgColor="DDE7F3")
    ws["A1"] = r.title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"{r.created_at:%d.%m.%Y %H:%M} — kayaAI"
    row = 4
    if r.summary:
        ws.cell(row, 1, "Özet").font = bold
        row += 1
        for para in _plain(r.summary).split("\n"):
            if para.strip():
                ws.cell(row, 1, para.strip()).alignment = Alignment(wrap_text=True, vertical="top")
                row += 1
        row += 1
    if r.key_figures:
        ws.cell(row, 1, "Anahtar rakamlar").font = bold
        row += 1
        for k, v in r.key_figures:
            ws.cell(row, 1, k)
            ws.cell(row, 2, v)
            row += 1
        row += 1
    for label, items in (("Notlar", r.notes), ("Kaynaklar", r.sources)):
        if items:
            ws.cell(row, 1, label).font = bold
            row += 1
            for n in items:
                ws.cell(row, 1, n)
                row += 1
            row += 1
    ws.column_dimensions["A"].width = 90
    ws.column_dimensions["B"].width = 28
    used: set[str] = {"Özet"}
    for i, t in enumerate(r.tables, 1):
        name = re.sub(r"[\[\]:*?/\\]", " ", t.title)[:28].strip() or f"Tablo {i}"
        while name in used:
            name = f"{name[:25]} {i}"
        used.add(name)
        sh = wb.create_sheet(name)
        sh.append(t.columns)
        for c in sh[1]:
            c.font, c.fill = bold, head_fill
        for data_row in t.rows:
            sh.append([v if isinstance(v, (int, float)) or v is None else str(v) for v in data_row])
        for j, col in enumerate(t.columns, 1):
            width = max([len(str(col))] + [len(_cell(rw[j - 1])) for rw in t.rows[:200] if j - 1 < len(rw)])
            sh.column_dimensions[get_column_letter(j)].width = min(60, max(10, width + 2))
            for cell in sh[get_column_letter(j)][1:]:
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
        sh.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_docx(r: Report) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(r.title, level=0)
    doc.add_paragraph(f"{r.created_at:%d.%m.%Y %H:%M} — kayaAI").runs[0].font.size = Pt(9)
    if r.summary:
        doc.add_heading("Özet", level=1)
        for para in r.summary.strip().split("\n"):
            p = para.strip()
            if not p:
                continue
            if re.match(r"^([-*•]|\d+[.)])\s+", p):
                doc.add_paragraph(_plain(re.sub(r"^([-*•]|\d+[.)])\s+", "", p)), style="List Bullet")
            else:
                doc.add_paragraph(_plain(p))
    if r.key_figures:
        doc.add_heading("Anahtar rakamlar", level=1)
        tbl = doc.add_table(rows=0, cols=2)
        tbl.style = "Light Grid Accent 1"
        for k, v in r.key_figures:
            cells = tbl.add_row().cells
            cells[0].text, cells[1].text = k, str(v)
    for t in r.tables:
        doc.add_heading(t.title, level=1)
        tbl = doc.add_table(rows=1, cols=len(t.columns))
        tbl.style = "Light Grid Accent 1"
        for j, c in enumerate(t.columns):
            tbl.rows[0].cells[j].text = str(c)
        for data_row in t.rows:
            cells = tbl.add_row().cells
            for j in range(len(t.columns)):
                cells[j].text = _cell(data_row[j]) if j < len(data_row) else ""
    for label, items in (("Notlar", r.notes), ("Kaynaklar", r.sources)):
        if items:
            doc.add_heading(label, level=1)
            for n in items:
                doc.add_paragraph(n, style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


FONT_CANDIDATES = [
    ("C:/Windows/Fonts/segoeui.ttf", "C:/Windows/Fonts/segoeuib.ttf"),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
]


def _register_font() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for regular, bold in FONT_CANDIDATES:
        if Path(regular).is_file():
            if "kayaai" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("kayaai", regular))
                pdfmetrics.registerFont(TTFont("kayaai-bold", bold if Path(bold).is_file() else regular))
            return "kayaai", "kayaai-bold"
    raise ReportError("PDF için Türkçe karakter destekleyen bir yazı tipi bulunamadı (Arial / Segoe UI / DejaVu Sans).")


def to_pdf(r: Report) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, bold = _register_font()
    base = ParagraphStyle("base", fontName=font, fontSize=10, leading=14)
    small = ParagraphStyle("small", parent=base, fontSize=8, textColor=colors.grey)
    h0 = ParagraphStyle("h0", parent=base, fontName=bold, fontSize=18, leading=22, spaceAfter=4)
    h1 = ParagraphStyle("h1", parent=base, fontName=bold, fontSize=13, leading=17, spaceBefore=10, spaceAfter=4)
    cellst = ParagraphStyle("cell", parent=base, fontSize=8.5, leading=11)

    def para(text: str, style: ParagraphStyle = base) -> Paragraph:
        t = escape(_plain(text))
        return Paragraph(t, style)

    def table(columns: list[str], rows: list[list[Any]]) -> Table:
        data = [[Paragraph(f"<b>{escape(str(c))}</b>", cellst) for c in columns]]
        data += [[Paragraph(escape(_cell(row[j]) if j < len(row) else ""), cellst) for j in range(len(columns))] for row in rows]
        tbl = Table(data, repeatRows=1, hAlign="LEFT")
        tbl.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE7F3")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA8B8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        return tbl

    story: list[Any] = [para(r.title, h0), para(f"{r.created_at:%d.%m.%Y %H:%M} — kayaAI", small), Spacer(1, 8)]
    if r.summary:
        story.append(para("Özet", h1))
        for p in r.summary.strip().split("\n"):
            if p.strip():
                story.append(para(("• " + re.sub(r"^([-*•]|\d+[.)])\s+", "", p.strip()))
                                  if re.match(r"^\s*([-*•]|\d+[.)])\s+", p) else p.strip()))
    if r.key_figures:
        story += [para("Anahtar rakamlar", h1), table(["", ""], [[k, v] for k, v in r.key_figures])]
    for t in r.tables:
        story += [para(t.title, h1), table(t.columns, t.rows[:500])]
    for label, items in (("Notlar", r.notes), ("Kaynaklar", r.sources)):
        if items:
            story.append(para(label, h1))
            story += [para("• " + n) for n in items]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=r.title, author="kayaAI", leftMargin=1.8 * cm, rightMargin=1.8 * cm,
                      topMargin=1.6 * cm, bottomMargin=1.6 * cm).build(story)
    return buf.getvalue()


WRITERS = {"md": to_markdown, "xlsx": to_xlsx, "docx": to_docx, "pdf": to_pdf}


def render(r: Report, fmt_: Format) -> bytes:
    if fmt_ not in WRITERS:
        raise ReportError(f"Bilinmeyen biçim: {fmt_}; biçimler: {', '.join(FORMATS)}")
    return WRITERS[fmt_](r)


def slug(text: str) -> str:
    tr = str.maketrans("çğıöşüÇĞİÖŞÜ", "cgiosuCGIOSU")
    s = re.sub(r"[^A-Za-z0-9]+", "-", text.translate(tr)).strip("-").lower()
    return s[:60] or "rapor"


def save(rt: "Runtime", r: Report, fmt_: Format, filename: str | None = None,
         requested_by: str = "user") -> "SubmitResult":
    """`~/kayaAI-files/raporlar/<ad>.<biçim>` — yeni dosya R1; aynı ad varsa R2 (onay)."""
    assert rt.files is not None
    data = render(r, fmt_)
    name = (filename or f"{slug(r.title)}-{r.created_at:%Y%m%d-%H%M%S}") + f".{fmt_}"
    return rt.files.write_bytes(f"{REPORT_DIR}/{Path(name).name}", data, description=f"rapor: {r.title}",
                                preview={"new_head": f"{r.title}\n\n{_plain(r.summary)[:300]}"},
                                requested_by=requested_by)
